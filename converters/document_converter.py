"""Document conversion using pypdf, python-docx, openpyxl, reportlab, markdown, html2text."""
import csv
import io
import os
import textwrap
from typing import List

# ── Conversion adjacency matrix ───────────────────────────────────────────────
CONVERSION_MAP = {
    "pdf":  ["txt"],
    "txt":  ["pdf", "docx", "html"],
    "docx": ["txt", "pdf", "html"],
    "html": ["txt", "pdf"],
    "htm":  ["txt", "pdf"],
    "csv":  ["xlsx"],
    "xlsx": ["csv"],
    "md":   ["html", "txt"],
}

SUPPORTED_INPUTS = set(CONVERSION_MAP.keys())


def get_document_targets(from_ext: str) -> List[str]:
    from_ext = from_ext.lower()
    return CONVERSION_MAP.get(from_ext, [])


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _read_text(path: str) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp949", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    return open(path, "r", errors="replace").read()


def _text_to_pdf(text: str, output_path: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))

    c = pdf_canvas.Canvas(output_path, pagesize=A4)
    width, height = A4
    margin_x, margin_y = 50, 50
    line_height = 16
    max_chars_per_line = 85
    y = height - margin_y

    c.setFont("HYGothic-Medium", 11)

    for raw_line in text.splitlines():
        # Wrap long lines
        wrapped = textwrap.wrap(raw_line, width=max_chars_per_line) or [""]
        for wline in wrapped:
            if y < margin_y + line_height:
                c.showPage()
                c.setFont("HYGothic-Medium", 11)
                y = height - margin_y
            c.drawString(margin_x, y, wline)
            y -= line_height

    c.save()


# ─────────────────────────────────────────────────────────────────────────────
#  PDF conversions
# ─────────────────────────────────────────────────────────────────────────────

def _pdf_to_txt(input_path: str, output_path: str) -> None:
    from pypdf import PdfReader
    reader = PdfReader(input_path)
    lines = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            lines.append(text)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(lines))


# ─────────────────────────────────────────────────────────────────────────────
#  TXT conversions
# ─────────────────────────────────────────────────────────────────────────────

def _txt_to_pdf(input_path: str, output_path: str) -> None:
    text = _read_text(input_path)
    _text_to_pdf(text, output_path)


def _txt_to_docx(input_path: str, output_path: str) -> None:
    from docx import Document
    text = _read_text(input_path)
    doc = Document()
    doc.add_heading("Converted Document", level=1)
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(output_path)


def _txt_to_html(input_path: str, output_path: str) -> None:
    text = _read_text(input_path)
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html = f"""<!DOCTYPE html>
<html lang="ko">
<head><meta charset="UTF-8"><title>Converted</title>
<style>body{{font-family:sans-serif;max-width:800px;margin:40px auto;line-height:1.6;white-space:pre-wrap;}}</style>
</head>
<body>{escaped}</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)


# ─────────────────────────────────────────────────────────────────────────────
#  DOCX conversions
# ─────────────────────────────────────────────────────────────────────────────

def _docx_to_txt(input_path: str, output_path: str) -> None:
    from docx import Document
    doc = Document(input_path)
    lines = [p.text for p in doc.paragraphs]
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def _docx_to_pdf(input_path: str, output_path: str) -> None:
    from docx import Document
    doc = Document(input_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    _text_to_pdf(text, output_path)


def _docx_to_html(input_path: str, output_path: str) -> None:
    from docx import Document
    doc = Document(input_path)
    paragraphs_html = []
    for p in doc.paragraphs:
        style = p.style.name if p.style else ""
        text = p.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if not text.strip():
            paragraphs_html.append("<br>")
        elif "Heading 1" in style:
            paragraphs_html.append(f"<h1>{text}</h1>")
        elif "Heading 2" in style:
            paragraphs_html.append(f"<h2>{text}</h2>")
        elif "Heading 3" in style:
            paragraphs_html.append(f"<h3>{text}</h3>")
        else:
            paragraphs_html.append(f"<p>{text}</p>")
    body = "\n".join(paragraphs_html)
    html = f"""<!DOCTYPE html>
<html lang="ko">
<head><meta charset="UTF-8"><title>Converted</title>
<style>body{{font-family:sans-serif;max-width:800px;margin:40px auto;line-height:1.7;}}</style>
</head>
<body>{body}</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)


# ─────────────────────────────────────────────────────────────────────────────
#  HTML conversions
# ─────────────────────────────────────────────────────────────────────────────

def _html_to_txt(input_path: str, output_path: str) -> None:
    import html2text
    text = _read_text(input_path)
    h = html2text.HTML2Text()
    h.ignore_links = False
    result = h.handle(text)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(result)


def _html_to_pdf(input_path: str, output_path: str) -> None:
    import html2text
    text_src = _read_text(input_path)
    h = html2text.HTML2Text()
    h.ignore_links = True
    plain = h.handle(text_src)
    _text_to_pdf(plain, output_path)


# ─────────────────────────────────────────────────────────────────────────────
#  CSV / XLSX conversions
# ─────────────────────────────────────────────────────────────────────────────

def _csv_to_xlsx(input_path: str, output_path: str) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    text = _read_text(input_path)
    reader = csv.reader(io.StringIO(text))
    for row in reader:
        ws.append(row)
    wb.save(output_path)


def _xlsx_to_csv(input_path: str, output_path: str) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(input_path)
    ws = wb.active
    with open(output_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([str(v) if v is not None else "" for v in row])


# ─────────────────────────────────────────────────────────────────────────────
#  Markdown conversions
# ─────────────────────────────────────────────────────────────────────────────

def _md_to_html(input_path: str, output_path: str) -> None:
    import markdown
    text = _read_text(input_path)
    body = markdown.markdown(text, extensions=["tables", "fenced_code"])
    html = f"""<!DOCTYPE html>
<html lang="ko">
<head><meta charset="UTF-8"><title>Converted</title>
<style>
body{{font-family:sans-serif;max-width:800px;margin:40px auto;line-height:1.7;color:#1a1a1a;}}
code{{background:#f4f4f4;padding:2px 6px;border-radius:4px;font-size:.9em;}}
pre{{background:#f4f4f4;padding:16px;border-radius:8px;overflow-x:auto;}}
table{{border-collapse:collapse;width:100%}}
th,td{{border:1px solid #ddd;padding:8px;text-align:left;}}
</style>
</head>
<body>{body}</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)


def _md_to_txt(input_path: str, output_path: str) -> None:
    import html2text, markdown
    text = _read_text(input_path)
    html_body = markdown.markdown(text)
    h = html2text.HTML2Text()
    h.ignore_links = False
    plain = h.handle(html_body)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(plain)


# ─────────────────────────────────────────────────────────────────────────────
#  Router
# ─────────────────────────────────────────────────────────────────────────────

_HANDLERS = {
    ("pdf",  "txt"):  _pdf_to_txt,
    ("txt",  "pdf"):  _txt_to_pdf,
    ("txt",  "docx"): _txt_to_docx,
    ("txt",  "html"): _txt_to_html,
    ("docx", "txt"):  _docx_to_txt,
    ("docx", "pdf"):  _docx_to_pdf,
    ("docx", "html"): _docx_to_html,
    ("html", "txt"):  _html_to_txt,
    ("html", "pdf"):  _html_to_pdf,
    ("htm",  "txt"):  _html_to_txt,
    ("htm",  "pdf"):  _html_to_pdf,
    ("csv",  "xlsx"): _csv_to_xlsx,
    ("xlsx", "csv"):  _xlsx_to_csv,
    ("md",   "html"): _md_to_html,
    ("md",   "txt"):  _md_to_txt,
}


def convert_document(input_path: str, output_path: str, output_format: str) -> None:
    from_ext = os.path.splitext(input_path)[1].lower().lstrip(".")
    to_ext = output_format.lower()
    handler = _HANDLERS.get((from_ext, to_ext))
    if not handler:
        raise ValueError(f"지원하지 않는 변환: {from_ext} → {to_ext}")
    handler(input_path, output_path)
